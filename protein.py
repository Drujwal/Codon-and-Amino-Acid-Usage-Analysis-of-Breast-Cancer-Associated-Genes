
from Bio import SeqIO
from Bio.Seq import Seq
from Bio.SeqUtils.ProtParam import ProteinAnalysis
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import os

# ============================================================
# INPUT AND OUTPUT FILES
# ============================================================

input_file = r"D:\Akshita\Paper\added analysis\protein\cancer.fa"

output_file = r"D:\Akshita\Paper\added analysis\protein\cancer_protein_analysis.xlsx"

# ============================================================
# CHECK INPUT FILE
# ============================================================

if not os.path.isfile(input_file):
    raise FileNotFoundError(f"FASTA file not found: {input_file}")

# ============================================================
# TRANSLATE NUCLEOTIDE SEQUENCES AND CALCULATE PROTEIN PROPERTIES
# ============================================================

results = []

for record in SeqIO.parse(input_file, "fasta"):

    gene_id = record.id
    dna_sequence = str(record.seq).upper().replace(" ", "")

    protein_length = None
    isoelectric_point = None
    molecular_weight = None
    status = "OK"

    if not dna_sequence:
        results.append([
            gene_id, 0, None, None, "Empty sequence"
        ])
        continue

    # Validate nucleotide characters
    if set(dna_sequence) - set("ACGTN"):
        results.append([
            gene_id, 0, None, None,
            "Invalid nucleotide characters"
        ])
        continue

    # CDS must be in complete codons
    if len(dna_sequence) % 3 != 0:
        results.append([
            gene_id, 0, None, None,
            "Length not divisible by 3; check CDS"
        ])
        continue

    # Translate using the standard genetic code
    protein_sequence = str(Seq(dna_sequence).translate())

    # Remove terminal stop codon
    if protein_sequence.endswith("*"):
        protein_sequence = protein_sequence[:-1]

    # Check for internal stop codons
    if "*" in protein_sequence:
        results.append([
            gene_id, len(protein_sequence), None, None,
            "Internal stop codon; check reading frame/CDS"
        ])
        continue

    # Check for ambiguous amino acids
    if "X" in protein_sequence:
        results.append([
            gene_id, len(protein_sequence), None, None,
            "Ambiguous amino acid; check nucleotide sequence"
        ])
        continue

    # Calculate protein properties
    analysis = ProteinAnalysis(protein_sequence)

    protein_length = len(protein_sequence)
    isoelectric_point = analysis.isoelectric_point()
    molecular_weight = analysis.molecular_weight()

    results.append([
        gene_id,
        protein_length,
        round(isoelectric_point, 6),
        round(molecular_weight, 2),
        status
    ])

# ============================================================
# CHECK FASTA CONTENT
# ============================================================

if not results:
    raise ValueError("No FASTA records were found in cancer.fa.")

# ============================================================
# CREATE EXCEL WORKBOOK
# ============================================================

wb = Workbook()
ws = wb.active
ws.title = "Protein Analysis"

headers = [
    "Gene_ID",
    "Protein_Length_aa",
    "Isoelectric_Point_pI",
    "Molecular_Weight_Da",
    "Status"
]

ws.append(headers)

for cell in ws[1]:
    cell.font = Font(bold=True)

for row in results:
    ws.append(row)

# ============================================================
# FORMAT EXCEL TABLE
# ============================================================

for row in range(2, len(results) + 2):
    ws.cell(row, 2).number_format = "0"
    ws.cell(row, 3).number_format = "0.000000"
    ws.cell(row, 4).number_format = "0.00"

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:E{len(results) + 1}"

column_widths = {
    "A": 20,
    "B": 22,
    "C": 25,
    "D": 25,
    "E": 55
}

for column, width in column_widths.items():
    ws.column_dimensions[column].width = width

# ============================================================
# ADD SUMMARY
# ============================================================

valid_count = sum(1 for row in results if row[4] == "OK")
review_count = len(results) - valid_count

summary_row = len(results) + 3

ws.cell(summary_row, 1, "Total FASTA records")
ws.cell(summary_row, 2, len(results))

ws.cell(summary_row + 1, 1, "Successfully analyzed proteins")
ws.cell(summary_row + 1, 2, valid_count)

ws.cell(summary_row + 2, 1, "Records requiring review")
ws.cell(summary_row + 2, 2, review_count)

for r in range(summary_row, summary_row + 3):
    ws.cell(r, 1).font = Font(bold=True)

# ============================================================
# SAVE EXCEL FILE
# ============================================================

wb.save(output_file)

print("=" * 55)
print("NUCLEOTIDE-TO-PROTEIN ANALYSIS COMPLETED")
print("=" * 55)
print(f"Total FASTA records: {len(results)}")
print(f"Successfully analyzed: {valid_count}")
print(f"Records requiring review: {review_count}")
print(f"Excel file: {output_file}")
print("=" * 55)